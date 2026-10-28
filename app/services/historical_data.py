"""
Loads real draw-history analysis from the workbook named in config.ini
(default: data/Mega_Millions_3_Year_Advanced_Pattern_Analysis.xlsx) and
exposes it as plain Python data the generation strategies can use.

If the file is missing or a sheet can't be read, we fall back to a small
built-in snapshot so the app still runs — callers don't need to special-case
"file not found".
"""
import threading
from openpyxl import load_workbook
from app.config import Config

# Fallback snapshot, used only if the workbook can't be read.
_FALLBACK_HOT = [42, 22, 49, 58, 66, 43, 27, 36, 52, 31]
_FALLBACK_OVERDUE = [11, 28, 69, 64, 15, 40, 12, 67, 16, 13]
_FALLBACK_PAIRS = [(22, 42), (18, 64), (18, 27), (43, 52), (19, 31)]


class HistoricalDataLoader:
    _lock = threading.Lock()
    _cache = None  # populated on first load()

    @classmethod
    def load(cls, force_reload: bool = False) -> dict:
        """Returns a dict with hot_numbers, overdue_numbers, top_pairs,
        mega_weights (ball -> weight for the current 1-24 regime), and
        a 'source' flag: 'workbook' or 'fallback'."""
        with cls._lock:
            if cls._cache is not None and not force_reload:
                return cls._cache
            cls._cache = cls._load_from_workbook()
            return cls._cache

    @classmethod
    def _load_from_workbook(cls) -> dict:
        path = Config.DATA_PATH
        if not path.exists():
            return cls._fallback(f"data file not found at {path}")

        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception as exc:  # corrupt/locked file, wrong format, etc.
            return cls._fallback(f"could not open workbook: {exc}")

        try:
            hot = cls._read_hot(wb)
            overdue = cls._read_overdue(wb)
            pairs = cls._read_pairs(wb)
            mega_weights = cls._read_mega_weights(wb)
        except Exception as exc:
            return cls._fallback(f"could not parse workbook: {exc}")
        finally:
            wb.close()

        if not hot or not overdue:
            return cls._fallback("workbook parsed but key sheets were empty")

        return {
            "hot_numbers": hot,
            "overdue_numbers": overdue,
            "top_pairs": pairs,
            "mega_weights": mega_weights,
            "source": "workbook",
            "path": str(path),
        }

    @staticmethod
    def _read_hot(wb) -> list:
        # "White Ball Frequency" sheet, columns: White Ball | Times Drawn | ...
        # Rows already come sorted by frequency (most-drawn first) in this report.
        ws = wb["White Ball Frequency"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        rows = [r for r in rows if r and r[0] is not None]
        rows.sort(key=lambda r: r[1] or 0, reverse=True)
        return [int(r[0]) for r in rows[: Config.HOT_COUNT]]

    @staticmethod
    def _read_overdue(wb) -> list:
        # "Overdue Numbers" sheet, columns: White Ball | Draws Since Last Seen | ...
        ws = wb["Overdue Numbers"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        rows = [r for r in rows if r and r[0] is not None]
        rows.sort(key=lambda r: r[1] or 0, reverse=True)
        return [int(r[0]) for r in rows[: Config.OVERDUE_COUNT]]

    @staticmethod
    def _read_pairs(wb) -> list:
        # "Top Pairs" sheet, columns: Ball A | Ball B | Times Together
        ws = wb["Top Pairs"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        rows = [r for r in rows if r and r[0] is not None and r[1] is not None]
        rows.sort(key=lambda r: r[2] or 0, reverse=True)
        return [(int(r[0]), int(r[1])) for r in rows[: Config.PAIR_COUNT]]

    @staticmethod
    def _read_mega_weights(wb) -> dict:
        # "Mega Ball Frequency" sheet, columns include "Current Regime (1-24)".
        # Weight every ball at least 1 so nothing is ever impossible to draw.
        ws = wb["Mega Ball Frequency"]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        try:
            col_idx = next(
                i for i, h in enumerate(header) if h and "Current Regime" in str(h)
            )
        except StopIteration:
            col_idx = 1  # fall back to "All 3 Years" column
        weights = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            ball = int(row[0])
            count = row[col_idx] if col_idx < len(row) else None
            weights[ball] = max(int(count or 0), 1)
        return weights

    @staticmethod
    def _fallback(reason: str) -> dict:
        return {
            "hot_numbers": _FALLBACK_HOT,
            "overdue_numbers": _FALLBACK_OVERDUE,
            "top_pairs": _FALLBACK_PAIRS,
            "mega_weights": {n: 1 for n in range(1, 25)},
            "source": "fallback",
            "reason": reason,
        }
